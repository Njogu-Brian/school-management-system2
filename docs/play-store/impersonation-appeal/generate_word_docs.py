"""Print-ready Google Play authorisation documents. No placeholders."""

from __future__ import annotations

from collections import deque
from pathlib import Path

import numpy as np
from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent
LOGO_SRC = Path(
    r"C:\Users\brian\.cursor\projects\d-Projects-school-management-system2-school-management-system2"
    r"\assets\c__Users_brian_AppData_Roaming_Cursor_User_workspaceStorage_"
    r"5a0fd611a9ffa73748e671d054ad43bd_images_print_logo-7832690e-1bce-4233-b29b-0fb58014c007.png"
)
LOGO_PNG = OUT / "royal-kings-crest.png"

BLACK = RGBColor(0x00, 0x00, 0x00)
FONT = "Times New Roman"

SCHOOL = "Royal Kings Premier School LTD"
DIRECTOR = "Purity Mwari Njogu"
PUBLISHER = "Brian Murage Njogu"
COMPANY = "Breysom Solutions"
DATE = "17th September 2026"
PACKAGE = "com.royalkingsschools.admin"
APP = "Royal Kings Admin"


def font(run, size=11, bold=False, italic=False, color=BLACK):
    run.font.name = FONT
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), FONT)
    rFonts.set(qn("w:hAnsi"), FONT)
    rFonts.set(qn("w:eastAsia"), FONT)
    rFonts.set(qn("w:cs"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def txt(p, text, **kwargs):
    run = p.add_run(text)
    font(run, **kwargs)
    return run


def space(p, before=0, after=6, align=None, justify=False, line=1.15):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align is not None:
        p.alignment = align


def no_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "FFFFFF")
        borders.append(el)
    tcPr.append(borders)


def cell_margins(cell, top=40, bottom=40, left=40, right=40):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def underline_para(p):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)


def prepare_logo() -> Path:
    im = Image.open(LOGO_SRC).convert("RGBA")
    arr = np.array(im)
    h, w = arr.shape[:2]
    dark = (arr[:, :, 0] < 40) & (arr[:, :, 1] < 40) & (arr[:, :, 2] < 40)
    seen = np.zeros((h, w), dtype=bool)
    q = deque()
    for x in range(w):
        q.append((0, x))
        q.append((h - 1, x))
    for y in range(h):
        q.append((y, 0))
        q.append((y, w - 1))
    while q:
        y, x = q.popleft()
        if y < 0 or x < 0 or y >= h or x >= w or seen[y, x]:
            continue
        seen[y, x] = True
        if not dark[y, x]:
            continue
        arr[y, x] = (255, 255, 255, 0)
        q.append((y + 1, x))
        q.append((y - 1, x))
        q.append((y, x + 1))
        q.append((y, x - 1))
    cut = Image.fromarray(arr)
    canvas = Image.new("RGB", cut.size, (255, 255, 255))
    canvas.paste(cut, mask=cut.split()[-1])
    canvas.save(LOGO_PNG, "PNG")
    return LOGO_PNG


def new_doc() -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(11)
    style.font.color.rgb = BLACK
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), FONT)
    rFonts.set(qn("w:hAnsi"), FONT)

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(2.0)

    footer = sec.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    txt(fp, f"{SCHOOL}, Wangige  ·  {DATE}", size=8, italic=True, color=RGBColor(0x44, 0x44, 0x44))
    return doc


def letterhead(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(3.6)
    table.columns[1].width = Cm(13.0)
    left, right = table.cell(0, 0), table.cell(0, 1)
    left.width = Cm(3.6)
    right.width = Cm(13.0)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    no_border(left)
    no_border(right)
    cell_margins(left, 0, 0, 0, 80)
    cell_margins(right, 40, 40, 80, 0)

    lp = left.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    space(lp, 0, 0)
    run = lp.add_run()
    font(run, size=11)
    run.add_picture(str(LOGO_PNG), height=Cm(2.4))

    p = right.paragraphs[0]
    space(p, 0, 0)
    txt(p, SCHOOL, size=14, bold=True)
    for line, sz, it in (
        ("Trading as Royal Kings School", 10, True),
        ("Wangige, Kiambu County, Kenya", 10, False),
        ("Email: info@royalkingsschools.sc.ke", 10, False),
        ("Website: www.royalkingsschools.sc.ke", 10, False),
    ):
        row = right.add_paragraph()
        space(row, 0, 0)
        txt(row, line, size=sz, italic=it)

    rule = doc.add_paragraph()
    space(rule, 8, 10)
    underline_para(rule)
    txt(rule, "")


def para(doc, text, *, after=8, before=0, bold=False, italic=False, justify=True):
    p = doc.add_paragraph()
    space(p, before, after, justify=justify)
    txt(p, text, bold=bold, italic=italic)
    return p


def heading(doc, text):
    p = doc.add_paragraph()
    space(p, 11, 4, justify=False)
    txt(p, text, size=11, bold=True)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    space(p, 0, 2, justify=False)
    txt(p, text, size=11)
    return p


def sign_block(cell, organisation_lines, name, role):
    no_border(cell)
    cell_margins(cell, 80, 80, 80, 120)
    first = True
    for line_text in organisation_lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        space(p, 0, 1)
        txt(p, line_text, size=11, bold=True)
    for _ in range(5):
        gap = cell.add_paragraph()
        space(gap, 8, 0)
        txt(gap, " ")
    line = cell.add_paragraph()
    space(line, 0, 2)
    txt(line, "........................................", size=11)
    for value in (name, role, f"Date: {DATE}"):
        if not value:
            continue
        row = cell.add_paragraph()
        space(row, 0, 1)
        txt(row, value, size=11)


def two_signatures(doc, left, right) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(8.3)
    table.columns[1].width = Cm(8.3)
    c0, c1 = table.cell(0, 0), table.cell(0, 1)
    c0.width = Cm(8.3)
    c1.width = Cm(8.3)
    sign_block(c0, *left)
    sign_block(c1, *right)


def build_letter() -> Path:
    doc = new_doc()
    doc.core_properties.title = "Authorisation letter — Royal Kings Admin"
    doc.core_properties.author = SCHOOL
    letterhead(doc)

    row = doc.add_paragraph()
    space(row, 0, 8, justify=False)
    txt(row, "Our Ref: RKPS/IT/GP/2026/09/17")
    tab = row.add_run("\t\t")
    font(tab)
    txt(row, f"Date: {DATE}")

    para(doc, "The Policy Team", after=0, justify=False)
    para(doc, "Google Play / Google LLC", after=10, justify=False)

    para(doc, "Dear Sir/Madam,", after=10, justify=False)

    re = doc.add_paragraph()
    space(re, 0, 10, justify=False)
    txt(re, "RE:  AUTHORISATION TO USE THE NAME AND LOGO OF ROYAL KINGS SCHOOL ON GOOGLE PLAY APPLICATION ", bold=True)
    txt(re, f"{APP} ({PACKAGE})", bold=True)

    para(
        doc,
        f"I, {DIRECTOR}, the Director of {SCHOOL}, a company incorporated in Kenya and trading as "
        f"Royal Kings School, of Wangige, Kiambu County, write to confirm the following.",
    )
    para(
        doc,
        f"{SCHOOL} is the owner of the names “Royal Kings”, “Royal Kings School”, “Royal Kings Premier School” "
        f"and “{SCHOOL}”, and of the school crest used on our website at https://royalkingsschools.sc.ke "
        f"(the purple circular mark with a crown, an open book and a pencil, and the words "
        f"“ROYAL KINGS SCHOOL — A Sure Foundation — Kindergarten | Primary | Junior Secondary”).",
    )
    para(
        doc,
        f"The School commissioned {COMPANY}, whose principal is {PUBLISHER}, to design, build and publish "
        f"a staff enterprise application for use only by authorised employees of the School. That application "
        f"is published on Google Play as follows:",
    )
    for line in (
        f"Application name: {APP}",
        f"Package name: {PACKAGE}",
        "Google Play Console developer account: Royal Kings Admin",
        f"Publisher: {COMPANY} / {PUBLISHER}",
        "Staff system: https://erp.royalkingsschools.sc.ke",
    ):
        bullet(doc, line)

    para(
        doc,
        f"{SCHOOL} hereby authorises {COMPANY} and {PUBLISHER} to use the School’s name, crest, colours, "
        f"tagline and contact details on that Google Play listing and in the application itself, including "
        f"the application name “{APP}”, the store icon, the feature graphic, screenshots, the short "
        f"description and the full description, and to state that the application is the official staff "
        f"application of Royal Kings School.",
        before=6,
    )
    para(
        doc,
        "This permission covers the store listing material identified in Google’s impersonation notice, namely "
        "the full description (en-GB), the application name (en-GB), the short description (en-GB), the "
        "application icon (en-GB) and the feature graphic (en-GB).",
    )
    para(
        doc,
        f"{PUBLISHER} and {COMPANY} are not impersonating the School. They are the School’s appointed "
        f"developer for this staff tool. The School supplied the name, logo and listing text. Access to the "
        f"application is limited to staff who hold school-issued credentials. It is not for students, parents "
        f"or the general public.",
    )
    para(
        doc,
        "This letter is given as signed written permission for Google Play policy review. It takes effect on "
        "the date below and continues until the School withdraws it in writing.",
    )
    para(doc, "Yours faithfully,", after=4, justify=False)

    two_signatures(
        doc,
        ((f"For {SCHOOL}",), DIRECTOR, "Director"),
        ((f"For {COMPANY}",), PUBLISHER, ""),
    )

    note = doc.add_paragraph()
    space(note, 16, 0, justify=False)
    txt(
        note,
        "Contact for verification: info@royalkingsschools.sc.ke  ·  https://royalkingsschools.sc.ke",
        size=10,
        italic=True,
    )

    path = OUT / "01-Authorization-Letter.docx"
    doc.save(path)
    return path


def build_licence() -> Path:
    doc = new_doc()
    doc.core_properties.title = "Brand licence — Royal Kings Admin"
    doc.core_properties.author = SCHOOL
    letterhead(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    space(title, 0, 2, align=WD_ALIGN_PARAGRAPH.CENTER)
    txt(title, "BRAND LICENCE AND DISTRIBUTION AGREEMENT", size=13, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    space(sub, 0, 12, align=WD_ALIGN_PARAGRAPH.CENTER)
    txt(sub, f"{APP}  ·  {PACKAGE}", size=11, italic=True)

    para(doc, f"This Agreement is made on {DATE}.", justify=False)
    heading(doc, "BETWEEN")
    para(
        doc,
        f"(1)  {SCHOOL}, trading as Royal Kings School, of Wangige, Kiambu County, Kenya, "
        f"email info@royalkingsschools.sc.ke (the “Licensor”); and",
    )
    para(
        doc,
        f"(2)  {COMPANY}, represented by {PUBLISHER} (the “Licensee”).",
    )
    para(doc, "The Licensor and the Licensee are together the “Parties”.")

    heading(doc, "BACKGROUND")
    para(
        doc,
        "A.  The Licensor owns the names, logo, crest and goodwill of Royal Kings School and Royal Kings "
        "Premier School, including the crest used at https://royalkingsschools.sc.ke.",
    )
    para(
        doc,
        "B.  The Licensor engaged the Licensee to build and publish a staff enterprise application for "
        "authorised employees of the School.",
    )
    para(
        doc,
        "C.  Google Play has asked for signed proof that the Licensee may use those brand assets on the "
        "store listing. This Agreement is that proof.",
    )

    heading(doc, "IT IS AGREED as follows")
    heading(doc, "1.  Licensed rights")
    para(
        doc,
        "1.1  The Licensor grants the Licensee a non-exclusive, royalty-free, non-transferable licence to "
        "use the Licensed Marks in connection with the App.",
    )
    para(doc, "1.2  “Licensed Marks” means:")
    for item in (
        "the words Royal Kings, Royal Kings School, Royal Kings Premier School, Royal Kings Premier School LTD, and Royal Kings Admin;",
        "the School crest and any version of it used as the Play Store icon;",
        "brand colours, the tagline “A Sure Foundation”, feature graphics and screenshots that include those marks;",
        "School contact details used in the store listing.",
    ):
        bullet(doc, item)
    para(
        doc,
        f"1.3  “App” means the Android application named {APP}, package {PACKAGE}, published from the "
        f"Google Play Console developer account “Royal Kings Admin”, including updates and testing-track builds.",
        before=6,
    )

    heading(doc, "2.  Scope of use")
    para(doc, "The Licensee may:")
    for item in (
        "publish and distribute the App on Google Play;",
        "use the Licensed Marks in the application name, short description, full description, icon, feature graphic, screenshots and in-app branding;",
        "state that the App is the official staff application of Royal Kings School for authorised employees;",
        "list the School website and official contact details in the store listing.",
    ):
        bullet(doc, item)

    heading(doc, "3.  Restrictions")
    para(
        doc,
        "3.1  This licence is only for the App in clause 1.3. It does not cover any other package name.",
    )
    para(
        doc,
        "3.2  The Licensee shall not present the App as being for students, parents or the general public, "
        "except to say that those groups may not use it.",
    )
    para(
        doc,
        "3.3  The Licensee shall not sub-licence the Licensed Marks except as required for distribution on Google Play.",
    )

    heading(doc, "4.  Ownership")
    para(
        doc,
        "Goodwill in the Licensed Marks belongs to the Licensor. This Agreement does not transfer ownership "
        "of the marks, the School’s data or the enterprise system.",
    )

    heading(doc, "5.  Term")
    para(
        doc,
        f"This Agreement starts on {DATE} and continues until the Licensor ends it by fourteen (14) days’ "
        f"written notice, or until the App is permanently unpublished. On ending, the Licensee shall remove "
        f"the Licensed Marks from Google Play within fourteen (14) days.",
    )

    heading(doc, "6.  Warranty")
    para(
        doc,
        f"The Licensor warrants that it owns or controls the Licensed Marks and has authority to grant this "
        f"licence. The Licensee warrants that it is the publisher of package {PACKAGE}.",
    )

    heading(doc, "7.  Governing law")
    para(doc, "This Agreement is governed by the laws of Kenya.")
    para(
        doc,
        "IN WITNESS WHEREOF the Parties have signed this Agreement on the date first written above.",
        before=8,
        italic=True,
        justify=False,
    )

    two_signatures(
        doc,
        (("For the Licensor", SCHOOL), DIRECTOR, "Director"),
        (("For the Licensee", COMPANY), PUBLISHER, ""),
    )

    path = OUT / "02-Brand-Licence-Agreement.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    prepare_logo()
    print(build_letter())
    print(build_licence())
